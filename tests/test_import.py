import io
from datetime import datetime

from docx import Document as DocxDocument
from openpyxl import Workbook

from app import db
from app.models import Contact, Election, Interaction
from app.services.importer import ImportError_, import_contacts

import pytest


def xlsx(rows):
    wb = Workbook()
    ws = wb.active
    ws.append(["Liste des sympathisants"])  # titre au-dessus des en-têtes
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


HEADER = ["Prénom", "NOM", "Adresse e-mail", "Téléphone", "Code postal", "Ville",
          "Étiquettes", "Commentaires", "Date de naissance"]


def jean(app):
    return Contact.query.filter_by(email="jean.dupont@example.org").one()


def test_excel_import_never_deletes(app):
    c = jean(app)
    c.tags, c.notes = "commerçant", "Tient la boulangerie."
    db.session.add(Interaction(contact=c, kind="appel", sentiment="positive"))
    db.session.commit()

    data = xlsx([
        HEADER,
        # Contact existant (e-mail en majuscules) : ville différente, téléphone vide.
        ["Jean", "Dupont", "JEAN.DUPONT@example.org", None, None, "Paris 11e",
         "bénévole", "Disponible le samedi", None],
        # Nouveau contact : Excel a supprimé les zéros initiaux.
        ["Lucie", "Bernard", None, 612345678, 1000, "Bourg-en-Bresse", None, None,
         datetime(1980, 5, 17)],
        [None, None, None, None, None, None, None, None, None],  # ligne vide
    ])
    report = import_contacts(data, "liste.xlsx")
    assert (report.created, report.updated) == (1, 1)

    c = jean(app)
    assert c.city == "Paris"                      # rien n'est écrasé par défaut
    assert c.phone == "06 12 34 56 78"            # cellule vide : conservé
    assert c.tags == "commerçant, bénévole"        # étiquettes ajoutées
    assert c.notes == "Tient la boulangerie.\nDisponible le samedi"
    assert len(c.interactions) == 1               # historique intact

    lucie = Contact.query.filter_by(last_name="Bernard").one()
    assert lucie.postal_code == "01000" and lucie.phone == "0612345678"
    assert lucie.birth_date.isoformat() == "1980-05-17"

    # Réimporter le même fichier ne crée aucun doublon.
    report = import_contacts(data, "liste.xlsx")
    assert report.created == 0 and report.unchanged == 2
    assert Contact.query.count() == 5


def test_overwrite_option(app):
    import_contacts(xlsx([HEADER, ["Jean", "Dupont", "jean.dupont@example.org", None, None,
                                   "Paris 11e", None, None, None]]), "a.xlsx", overwrite=True)
    c = jean(app)
    assert c.city == "Paris 11e" and c.phone == "06 12 34 56 78"


def test_match_by_name_without_email(app):
    data = xlsx([["prenom", "nom", "ville"], ["jean", "DUPONT", None]])
    assert import_contacts(data, "b.xlsx").created == 0


def test_vote_column(app):
    db.session.add(Election(name="Municipales", is_current=True))
    db.session.commit()
    import_contacts(xlsx([["Nom", "Prénom", "A voté"], ["Martin", "Pierre", "oui"]]), "v.xlsx")
    election = Election.current()
    assert Contact.query.filter_by(last_name="Martin").one().vote_for(election).voted is True


def test_word_table(app):
    doc = DocxDocument()
    doc.add_paragraph("Contacts du marché")
    table = doc.add_table(rows=2, cols=3)
    for cell, text in zip(table.rows[0].cells, ["Nom", "Prénom", "Ville"]):
        cell.text = text
    for cell, text in zip(table.rows[1].cells, ["Garcia", "Inès", "Toulouse"]):
        cell.text = text
    buf = io.BytesIO()
    doc.save(buf)
    assert import_contacts(buf.getvalue(), "liste.docx").created == 1
    assert Contact.query.filter_by(last_name="Garcia").one().city == "Toulouse"


def test_errors(app):
    with pytest.raises(ImportError_):
        import_contacts(b"x", "vieux.xls")
    with pytest.raises(ImportError_):
        import_contacts(xlsx([["Ville"], ["Lyon"]]), "sans_nom.xlsx")
