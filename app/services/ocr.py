"""Extraction du texte d'un document (photo/scan, PDF, Word ou texte)."""
import io
import logging
import shutil

log = logging.getLogger(__name__)

IMAGE_EXTENSIONS = {"png", "jpg", "jpeg", "tif", "tiff", "bmp", "gif", "webp"}
ALLOWED_EXTENSIONS = IMAGE_EXTENSIONS | {"pdf", "txt", "docx"}


class OCRUnavailable(RuntimeError):
    pass


def ocr_available():
    return shutil.which("tesseract") is not None


def _ocr_image(image, lang):
    if not ocr_available():
        raise OCRUnavailable(
            "Tesseract n'est pas installé sur le serveur (paquets tesseract-ocr et tesseract-ocr-fra)."
        )
    import pytesseract
    from PIL import ImageOps

    image = ImageOps.exif_transpose(image)
    if image.mode not in ("L", "RGB"):
        image = image.convert("RGB")
    # Les petites images sont agrandies : Tesseract lit mieux vers 300 dpi.
    if image.width < 1500:
        factor = 1500 / image.width
        image = image.resize((int(image.width * factor), int(image.height * factor)))
    try:
        return pytesseract.image_to_string(image.convert("L"), lang=lang)
    except pytesseract.TesseractError:
        # Langue non installée : on retente avec l'anglais seul.
        return pytesseract.image_to_string(image.convert("L"))


def _extract_pdf(data, lang):
    from PIL import Image
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    parts = []
    for page in reader.pages:
        text = (page.extract_text() or "").strip()
        if len(text) >= 20:
            parts.append(text)
            continue
        # Page scannée : on passe les images de la page à l'OCR.
        for img in page.images:
            try:
                parts.append(_ocr_image(Image.open(io.BytesIO(img.data)), lang))
            except OCRUnavailable:
                raise
            except Exception as exc:  # image illisible : on continue
                log.warning("Image de PDF ignorée : %s", exc)
    return "\n".join(parts)


def _extract_docx(data, lang):
    from docx import Document as DocxDocument
    from PIL import Image

    doc = DocxDocument(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" ".join(cell.text.strip() for cell in row.cells))
    # En-têtes / pieds de page : souvent l'expéditeur y figure.
    for section in doc.sections:
        for part in (section.header, section.footer):
            parts.extend(p.text for p in part.paragraphs if p.text.strip())
    # Images collées dans le document (courrier scanné inséré dans Word).
    if ocr_available():
        for rel in doc.part.rels.values():
            if "image" in rel.reltype:
                try:
                    parts.append(_ocr_image(Image.open(io.BytesIO(rel.target_part.blob)), lang))
                except Exception as exc:
                    log.warning("Image Word ignorée : %s", exc)
    return "\n".join(parts)


def extract_text(data, filename, lang="fra+eng"):
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext == "txt":
        for encoding in ("utf-8", "latin-1"):
            try:
                return data.decode(encoding)
            except UnicodeDecodeError:
                continue
    if ext == "pdf":
        return _extract_pdf(data, lang)
    if ext == "docx":
        return _extract_docx(data, lang)
    if ext in IMAGE_EXTENSIONS:
        from PIL import Image

        image = Image.open(io.BytesIO(data))
        pages = []
        # TIFF multi-pages (fréquent avec les scanners)
        for frame in range(getattr(image, "n_frames", 1)):
            image.seek(frame)
            pages.append(_ocr_image(image.copy(), lang))
        return "\n".join(pages)
    raise ValueError(f"Format non pris en charge : .{ext}")
